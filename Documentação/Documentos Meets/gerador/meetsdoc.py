"""Gerador dos documentos Word do Meets a partir do modelo "Documento Visão Meets.docx".

Mantém capa, cabeçalho, rodapé, estilos e numeração de títulos do modelo da equipe e
substitui o corpo. O índice é escrito com o número de página real: o documento é
gerado, convertido para PDF (LibreOffice) e gerado de novo com as páginas encontradas.

Uso: python3 gerar.py   (na pasta gerador/)
"""
import copy
import os
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(AQUI)
MODELO = os.path.join(AQUI, 'modelo.docx')  # Documento de Visão original da equipe
DIAGRAMAS = os.path.join(DOCS, 'diagramas_LES')

BULLET_NUM_ID = 9          # ● ○ ■ do modelo
DECIMAL_ABSTRACT_ID = 13   # 1) a) i) do modelo
COR_CABECALHO = 'D9E2F3'


def _markup_runs(paragraph, text, size=None, italic=False):
    """Escreve texto com **negrito** inline."""
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        run = paragraph.add_run(part[2:-2] if bold else part)
        run.bold = bold or None
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)
    return paragraph


class MeetsDoc:
    def __init__(self, titulo, saida):
        self.titulo = titulo
        self.saida = saida
        self.doc = Document(MODELO)
        self._limpar_corpo()
        self._ajustar_capa_e_cabecalho()
        self.headings = []      # (nivel, texto)
        self.toc_paragraphs = []
        self._next_num_id = None

    # ---------- estrutura do modelo ----------
    def _limpar_corpo(self):
        body = self.doc.element.body
        filhos = list(body)
        # capa = parágrafos até o título "Documento de Visão" (+ espaçamento)
        corte = None
        for i, p in enumerate(self.doc.paragraphs):
            if p.style.name == 'Heading 1':
                corte = p._p
                break
        manter = True
        for el in filhos:
            if el is corte:
                manter = False
            if el.tag == qn('w:sectPr'):
                continue
            if not manter:
                body.remove(el)
        # ponto de inserção: antes do sectPr final
        self._sectPr = body.find(qn('w:sectPr'))

    def _ajustar_capa_e_cabecalho(self):
        for p in self.doc.paragraphs:
            if p.text.strip() == 'Documento de Visão':
                for r in p.runs[1:]:
                    r.text = ''
                p.runs[0].text = self.titulo
        for part in self.doc.part.package.parts:
            if part.partname.startswith('/word/header'):
                xml = part.blob.decode('utf-8') if hasattr(part, 'blob') else None
        for section in self.doc.sections:
            for hdr in (section.header, section.first_page_header, section.even_page_header):
                for p in hdr.paragraphs:
                    if 'Documento de Visão' not in p.text:
                        continue
                    runs = p.runs
                    for i, r in enumerate(runs):
                        if r.text.startswith('Documento'):
                            r.text = self.titulo
                            for resto in runs[i + 1:]:
                                resto.text = ''
                            break

    def _append(self, element):
        self._sectPr.addprevious(element)
        return element

    def _new_paragraph(self, style=None):
        p = self.doc.add_paragraph(style=style)
        el = p._p
        el.getparent().remove(el)
        self._append(el)
        return p

    # ---------- blocos ----------
    def indice(self):
        h = self.h1('Índice', quebra=False)
        h.paragraph_format.page_break_before = True
        self._toc_anchor = self._new_paragraph()

    def h1(self, texto, quebra=False):
        if quebra and self.headings:
            self.page_break()
        p = self._new_paragraph('Heading 1')
        p.add_run(texto)
        self.headings.append((1, texto))
        return p

    def h2(self, texto):
        p = self._new_paragraph('Heading 2')
        p.add_run(texto)
        self.headings.append((2, texto))
        return p

    def h3(self, texto):
        p = self._new_paragraph('Heading 3')
        p.add_run(texto)
        return p

    def p(self, texto, italic=False, align=None, size=None, space_after=6):
        par = self._new_paragraph()
        _markup_runs(par, texto, size=size, italic=italic)
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.line_spacing = 1.15
        if align == 'center':
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align != 'left':
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return par

    def rotulo(self, texto):
        """Rótulo em negrito (ex.: 'Pré-condições:')."""
        par = self._new_paragraph()
        r = par.add_run(texto)
        r.bold = True
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.keep_with_next = True
        return par

    def _numpr(self, par, num_id, level):
        pPr = par._p.get_or_add_pPr()
        numPr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), str(level))
        nid = OxmlElement('w:numId'); nid.set(qn('w:val'), str(num_id))
        numPr.append(ilvl); numPr.append(nid)
        pPr.append(numPr)

    def bullets(self, itens, level=0):
        """itens: str ou (str, [subitens])."""
        for item in itens:
            sub = None
            if isinstance(item, tuple):
                item, sub = item
            par = self._new_paragraph()
            _markup_runs(par, item)
            self._numpr(par, BULLET_NUM_ID, level)
            par.paragraph_format.space_after = Pt(2)
            if sub:
                self.bullets(sub, level + 1)

    def _novo_num_decimal(self):
        numbering = self.doc.part.numbering_part.element
        ids = [int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))]
        new_id = max(ids) + 1
        num = OxmlElement('w:num'); num.set(qn('w:numId'), str(new_id))
        absid = OxmlElement('w:abstractNumId'); absid.set(qn('w:val'), str(DECIMAL_ABSTRACT_ID))
        num.append(absid)
        ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), '0')
        st = OxmlElement('w:startOverride'); st.set(qn('w:val'), '1')
        ov.append(st); num.append(ov)
        numbering.append(num)
        return new_id

    def numerada(self, itens):
        num_id = self._novo_num_decimal()
        for item in itens:
            par = self._new_paragraph()
            _markup_runs(par, item)
            self._numpr(par, num_id, 0)
            par.paragraph_format.space_after = Pt(2)

    def tabela(self, cabecalho, linhas, larguras_cm=None, fonte=9):
        tbl = self.doc.add_table(rows=1, cols=len(cabecalho))
        el = tbl._tbl
        el.getparent().remove(el)
        self._append(el)
        tbl.style = self.doc.styles['Table Grid']
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        def preencher(cell, texto, bold=False, shade=None):
            cell.text = ''
            par = cell.paragraphs[0]
            linhas_txt = str(texto).split('\n')
            for i, t in enumerate(linhas_txt):
                if i:
                    par = cell.add_paragraph()
                _markup_runs(par, t, size=fonte)
                if bold:
                    for r in par.runs:
                        r.bold = True
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.0
                # tamanho da marca de parágrafo (células vazias não ficam altas)
                pPr = par._p.get_or_add_pPr()
                rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(fonte * 2)); rpr.append(sz); pPr.append(rpr)
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if shade:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), shade)
                tcPr.append(shd)

        for i, h in enumerate(cabecalho):
            preencher(tbl.rows[0].cells[i], h, bold=True, shade=COR_CABECALHO)
        # repetir cabeçalho em cada página
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
        for linha in linhas:
            row = tbl.add_row()
            cells = row.cells
            if isinstance(linha, str):
                # linha de grupo: ocupa a tabela inteira
                merged = cells[0].merge(cells[-1])
                preencher(merged, linha, bold=True, shade='F2F2F2')
                continue
            for i, v in enumerate(linha):
                preencher(cells[i], v)
        if larguras_cm:
            for row in tbl.rows:
                for i, w in enumerate(larguras_cm):
                    row.cells[i].width = Cm(w)
            grid = el.tblGrid
            for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
                gc.set(qn('w:w'), str(int(larguras_cm[i] * 567)))
            tblPr0 = el.tblPr
            for old in tblPr0.findall(qn('w:tblW')):
                tblPr0.remove(old)
            tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(larguras_cm) * 567))); tw.set(qn('w:type'), 'dxa')
            tblPr0.append(tw)
            lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr0.append(lay)
        # margem interna das células
        tblPr = el.tblPr
        mar = OxmlElement('w:tblCellMar')
        for lado in ('top', 'bottom'):
            e = OxmlElement(f'w:{lado}'); e.set(qn('w:w'), '40'); e.set(qn('w:type'), 'dxa'); mar.append(e)
        for lado in ('left', 'right'):
            e = OxmlElement(f'w:{lado}'); e.set(qn('w:w'), '80'); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblPr.append(mar)
        self._new_paragraph().paragraph_format.space_after = Pt(4)
        return tbl

    def imagem(self, nome, largura_cm=17.0, legenda=None, altura_max_cm=21.0):
        caminho = os.path.join(DIAGRAMAS, nome)
        from PIL import Image
        with Image.open(caminho) as im:
            w, h = im.size
        largura = largura_cm
        if h / w * largura > altura_max_cm:
            largura = altura_max_cm * w / h
        par = self._new_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(caminho, width=Cm(largura))
        if legenda:
            leg = self._new_paragraph()
            leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = leg.add_run(legenda); r.italic = True; r.font.size = Pt(9)
            leg.paragraph_format.space_after = Pt(10)

    def page_break(self):
        par = self._new_paragraph()
        par.add_run().add_break(WD_BREAK.PAGE)

    # ---------- índice e saída ----------
    def _escrever_indice(self, paginas):
        anchor = self._toc_anchor._p
        contadores = [0, 0]
        for (nivel, texto) in self.headings:
            if nivel == 1:
                contadores[0] += 1; contadores[1] = 0
                num = f'{contadores[0]}.'
            else:
                contadores[1] += 1
                num = f'{contadores[0]}.{contadores[1]}'
            pg = paginas.get((nivel, texto), '')
            par = self.doc.add_paragraph(style='toc 1' if nivel == 1 else 'toc 2')
            par.add_run(f'{num}\t{texto if pg else "§" + texto}\t{pg}')
            from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
            ts = par.paragraph_format.tab_stops
            ts.add_tab_stop(Cm(1.0 if nivel == 1 else 1.6))
            ts.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if nivel == 2:
                par.paragraph_format.left_indent = Cm(0.6)
            el = par._p; el.getparent().remove(el)
            anchor.addprevious(el)

    def _paginas_pdf(self, docx_path):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, docx_path],
                           check=True, capture_output=True)
            pdf = os.path.join(tmp, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf')
            txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], check=True, capture_output=True).stdout.decode('utf-8')
        paginas_txt = txt.split('\f')
        resultado = {}
        contadores = [0, 0]
        inicio = 1  # pula a capa
        for i, pg in enumerate(paginas_txt):
            if re.search(r'1\.\s+Índice', pg):
                inicio = i  # as linhas do índice levam '§' e não casam com os títulos
                break
        for (nivel, texto) in self.headings:
            if nivel == 1:
                contadores[0] += 1; contadores[1] = 0
                rot = f'{contadores[0]}.'
            else:
                contadores[1] += 1
                rot = f'{contadores[0]}.{contadores[1]}'
            alvo = re.sub(r'\s+', ' ', f'{rot} {texto}').strip()
            if texto == 'Índice':
                resultado[(nivel, texto)] = inicio + 1
                continue
            for i in range(inicio, len(paginas_txt)):
                pagina = re.sub(r'\s+', ' ', paginas_txt[i])
                if alvo in pagina:
                    resultado[(nivel, texto)] = i + 1
                    inicio = i
                    break
        return resultado

    def salvar(self):
        # passo 1: índice sem páginas, para medir
        estado = copy.deepcopy(self.doc.element.body)
        self._escrever_indice({})  # sem páginas: textos marcados com '§'
        self.doc.save(self.saida)
        paginas = self._paginas_pdf(self.saida)
        # passo 2: restaura e escreve com páginas
        body = self.doc.element.body
        for el in list(body):
            body.remove(el)
        for el in estado:
            body.append(el)
        self._sectPr = body.find(qn('w:sectPr'))
        # reencontra a âncora do índice (parágrafo vazio após o título Índice)
        for i, p in enumerate(self.doc.paragraphs):
            if p.style.name == 'Heading 1' and p.text == 'Índice':
                self._toc_anchor = self.doc.paragraphs[i + 1]
                break
        self._escrever_indice(paginas)
        self.doc.save(self.saida)
        faltando = [t for (n, t) in self.headings if (n, t) not in paginas]
        return faltando
